"""
Tests for the document markup engine.

Creates a sample .docx, applies changes, and verifies the output.
"""

import os
import tempfile

from docx import Document
from docx.shared import RGBColor

from app.services.document_markup import (
    Change,
    DocumentMarkupEngine,
    extract_document_text,
    extract_document_structure,
    DELETION_COLOR,
    INSERTION_COLOR,
)


def _create_sample_nda(path: str):
    """Create a sample NDA document for testing."""
    doc = Document()

    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=0)

    doc.add_paragraph(
        "This Non-Disclosure Agreement (the \"Agreement\") is entered into as of "
        "January 1, 2025 (the \"Effective Date\") by and between:"
    )

    doc.add_paragraph(
        "Company ABC, Inc., a Delaware corporation (\"Disclosing Party\"), and "
        "XYZ Corp., a California corporation (\"Receiving Party\")."
    )

    doc.add_heading("1. Definition of Confidential Information", level=1)
    doc.add_paragraph(
        "\"Confidential Information\" means any and all information or data, "
        "whether oral, written, electronic, or visual, disclosed by the Disclosing "
        "Party to the Receiving Party, including but not limited to trade secrets, "
        "business plans, financial information, customer lists, and technical data."
    )

    doc.add_heading("2. Obligations of Receiving Party", level=1)
    doc.add_paragraph(
        "The Receiving Party shall hold and maintain the Confidential Information "
        "in strict confidence for the sole and exclusive benefit of the Disclosing "
        "Party. The Receiving Party shall not, without the prior written approval "
        "of the Disclosing Party, use for the Receiving Party's own benefit, "
        "publish, copy, or otherwise disclose to others, or permit the use by "
        "others for their benefit or to the detriment of the Disclosing Party."
    )

    doc.add_heading("3. Term", level=1)
    doc.add_paragraph(
        "This Agreement shall remain in effect for a period of five (5) years "
        "from the Effective Date, unless earlier terminated by either party "
        "upon thirty (30) days written notice."
    )

    doc.add_heading("4. Remedies", level=1)
    doc.add_paragraph(
        "The Receiving Party acknowledges that any breach of this Agreement "
        "may cause irreparable harm to the Disclosing Party. The Disclosing "
        "Party shall be entitled to seek injunctive relief and any other "
        "remedies available at law or in equity."
    )

    doc.add_heading("5. Governing Law", level=1)
    doc.add_paragraph(
        "This Agreement shall be governed by and construed in accordance with "
        "the laws of the State of Delaware, without regard to its conflict of "
        "law provisions."
    )

    # Add a table for testing table handling
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Party"
    table.cell(0, 1).text = "Signature"
    table.cell(1, 0).text = "Disclosing Party: Company ABC, Inc."
    table.cell(1, 1).text = "________________________"

    doc.save(path)


def test_extract_text():
    """Test text extraction from a Word document."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name

    try:
        _create_sample_nda(path)
        text = extract_document_text(path)

        assert "Non-Disclosure Agreement" in text
        assert "Confidential Information" in text
        assert "Receiving Party" in text
        assert "five (5) years" in text
        print("PASS: test_extract_text")
    finally:
        os.unlink(path)


def test_extract_structure():
    """Test document structure extraction."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        path = f.name

    try:
        _create_sample_nda(path)
        structure = extract_document_structure(path)

        assert len(structure) > 0
        # Check that headings are identified
        headings = [s for s in structure if s["is_heading"]]
        assert len(headings) >= 5  # Our sample has 5 section headings
        print("PASS: test_extract_structure")
    finally:
        os.unlink(path)


def test_deletion():
    """Test applying a deletion (red strikethrough)."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="deletion",
                original_text="including but not limited to",
                rationale="Overly broad catch-all language",
            )
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        assert stats["applied"] == 1
        assert stats["failed"] == 0

        # Verify the output document
        doc = Document(out_path)
        found_strikethrough = False
        for para in doc.paragraphs:
            for run in para.runs:
                if run.font.strike and run.font.color.rgb == DELETION_COLOR:
                    if "including but not limited to" in run.text:
                        found_strikethrough = True

        assert found_strikethrough, "Deletion formatting not found in output"
        print("PASS: test_deletion")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


def test_replacement():
    """Test applying a replacement (red strikethrough + blue underline)."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="replacement",
                original_text="five (5) years",
                new_text="two (2) years",
                rationale="Standard NDA term should be 2 years, not 5",
                section="3. Term",
                priority="high",
            )
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        assert stats["applied"] == 1

        # Verify both deletion and insertion formatting
        doc = Document(out_path)
        found_deletion = False
        found_insertion = False

        for para in doc.paragraphs:
            for run in para.runs:
                if run.font.strike and "five (5) years" in run.text:
                    found_deletion = True
                if run.font.underline and "two (2) years" in run.text:
                    found_insertion = True

        assert found_deletion, "Replacement deletion formatting not found"
        assert found_insertion, "Replacement insertion formatting not found"
        print("PASS: test_replacement")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


def test_insertion():
    """Test inserting new text after anchor text."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="insertion",
                original_text="thirty (30) days written notice",
                new_text=" to the other party's registered address",
                rationale="Specify how notice should be delivered",
                section="3. Term",
            )
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        assert stats["applied"] == 1

        # Verify insertion formatting exists
        doc = Document(out_path)
        found_insertion = False
        for para in doc.paragraphs:
            for run in para.runs:
                if run.font.underline and "registered address" in run.text:
                    found_insertion = True

        assert found_insertion, "Insertion formatting not found"
        print("PASS: test_insertion")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


def test_multiple_changes():
    """Test applying multiple changes to the same document."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="replacement",
                original_text="five (5) years",
                new_text="two (2) years",
                rationale="Reduce term length",
                priority="high",
            ),
            Change(
                change_type="deletion",
                original_text="including but not limited to",
                rationale="Remove catch-all",
                priority="medium",
            ),
            Change(
                change_type="replacement",
                original_text="State of Delaware",
                new_text="State of California",
                rationale="Change governing law to match receiving party jurisdiction",
                priority="medium",
            ),
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        assert stats["applied"] == 3
        assert stats["failed"] == 0
        print("PASS: test_multiple_changes")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


def test_comment_only():
    """Test adding a comment without modifying text."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="comment_only",
                original_text="irreparable harm",
                rationale="Consider whether this standard is appropriate for the type of information being shared",
            )
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        assert stats["applied"] == 1
        # The document should open without errors
        doc = Document(out_path)
        assert len(doc.paragraphs) > 0
        print("PASS: test_comment_only")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


def test_unfound_text_fails_gracefully():
    """Test that text not found in the document is handled gracefully."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="deletion",
                original_text="this text does not exist in the document at all",
                rationale="Should fail gracefully",
            )
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        assert stats["applied"] == 0
        assert stats["failed"] == 1
        print("PASS: test_unfound_text_fails_gracefully")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


def test_document_opens_cleanly():
    """Test that the generated document can be opened and read back."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        src_path = f.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        out_path = f.name

    try:
        _create_sample_nda(src_path)

        changes = [
            Change(
                change_type="replacement",
                original_text="five (5) years",
                new_text="two (2) years",
                rationale="Reduce term",
            ),
            Change(
                change_type="deletion",
                original_text="including but not limited to",
                rationale="Remove overbroad language",
            ),
            Change(
                change_type="insertion",
                original_text="thirty (30) days written notice",
                new_text=" via certified mail",
                rationale="Specify delivery method",
            ),
            Change(
                change_type="comment_only",
                original_text="irreparable harm",
                rationale="Review this standard",
            ),
        ]

        engine = DocumentMarkupEngine()
        stats = engine.apply_markup(src_path, changes, out_path)

        # Verify we can open the result
        result_doc = Document(out_path)
        assert len(result_doc.paragraphs) > 0

        # Verify the original structure is preserved
        original_doc = Document(src_path)
        # Paragraph count may differ slightly due to splits, but should be close
        assert abs(len(result_doc.paragraphs) - len(original_doc.paragraphs)) <= 2

        print("PASS: test_document_opens_cleanly")
        print(f"Stats: {stats}")
    finally:
        os.unlink(src_path)
        os.unlink(out_path)


if __name__ == "__main__":
    test_extract_text()
    test_extract_structure()
    test_deletion()
    test_replacement()
    test_insertion()
    test_multiple_changes()
    test_comment_only()
    test_unfound_text_fails_gracefully()
    test_document_opens_cleanly()
    print("\nAll tests passed!")
