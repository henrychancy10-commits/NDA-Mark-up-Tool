"""
Integration test: creates sample files, uploads via Flask API, and tests
the full document markup pipeline (skipping Claude API call, using mock changes).
"""

import io
import json
import os
import sys
import tempfile

from docx import Document

# Add parent directory to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.create_app import create_app
from app.services.document_markup import Change, DocumentMarkupEngine


def create_sample_nda_bytes():
    """Create a sample NDA as in-memory bytes."""
    doc = Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=0)
    doc.add_paragraph(
        "This Non-Disclosure Agreement (the \"Agreement\") is entered into as of "
        "January 1, 2025 (the \"Effective Date\") by and between:"
    )
    doc.add_paragraph(
        "Company ABC, Inc., a Delaware corporation (\"Disclosing Party\"), and "
        "XYZ Corp., a California corporation (\"Receiving Party\")."
    )
    doc.add_heading("1. Definition of Confidential Information", level=1)
    doc.add_paragraph(
        "\"Confidential Information\" means any and all information or data, "
        "whether oral, written, electronic, or visual, disclosed by the Disclosing "
        "Party to the Receiving Party, including but not limited to trade secrets, "
        "business plans, financial information, customer lists, and technical data."
    )
    doc.add_heading("2. Term", level=1)
    doc.add_paragraph(
        "This Agreement shall remain in effect for a period of five (5) years "
        "from the Effective Date, unless earlier terminated by either party "
        "upon thirty (30) days written notice."
    )
    doc.add_heading("3. Governing Law", level=1)
    doc.add_paragraph(
        "This Agreement shall be governed by and construed in accordance with "
        "the laws of the State of Delaware, without regard to its conflict of "
        "law provisions."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def create_sample_guidelines_bytes():
    """Create sample markup guidelines as text bytes."""
    content = """NDA MARKUP GUIDELINES

1. Term Length: NDA term should not exceed 2 years. Replace any term longer than 2 years.
2. Overbroad Language: Remove "including but not limited to" catch-all phrases.
3. Governing Law: Prefer California law for receiving party protection.
4. Notice Period: 30-day notice period is acceptable.
5. Confidential Information: Definition should be specific, not overly broad.
"""
    return io.BytesIO(content.encode("utf-8"))


def test_full_pipeline():
    """Test the complete flow: upload, mock-analyze, review, generate."""
    app = create_app()

    with app.test_client() as client:
        # 1. Create project
        r = client.post(
            "/api/projects",
            json={"name": "Integration Test", "description": "Automated test"},
        )
        assert r.status_code == 201
        project = r.get_json()
        project_id = project["id"]
        print(f"Created project: {project_id}")

        # 2. Upload guidelines
        guidelines_data = create_sample_guidelines_bytes()
        r = client.post(
            f"/api/projects/{project_id}/guidelines",
            data={"file": (guidelines_data, "guidelines.txt")},
            content_type="multipart/form-data",
        )
        assert r.status_code == 201
        print(f"Uploaded guidelines: {r.get_json()}")

        # 3. Upload NDA document
        nda_data = create_sample_nda_bytes()
        r = client.post(
            f"/api/projects/{project_id}/documents",
            data={
                "file": (nda_data, "test_nda.docx"),
                "mode": "balanced",
            },
            content_type="multipart/form-data",
        )
        assert r.status_code == 201
        doc = r.get_json()
        doc_id = doc["id"]
        print(f"Uploaded document: {doc_id}")

        # 4. Get document status
        r = client.get(f"/api/documents/{doc_id}")
        assert r.status_code == 200
        doc_data = r.get_json()
        assert doc_data["status"] == "uploaded"
        print(f"Document status: {doc_data['status']}")

        # 5. Apply mock changes directly (bypassing Claude API)
        # This simulates what the analyze endpoint would do
        from app.models.database import save_changes, update_document_status

        mock_changes = [
            {
                "change_type": "replacement",
                "original_text": "five (5) years",
                "new_text": "two (2) years",
                "rationale": "NDA term should not exceed 2 years per guidelines",
                "section": "2. Term",
                "priority": "high",
                "status": "pending",
                "applied": 0,
            },
            {
                "change_type": "deletion",
                "original_text": "including but not limited to",
                "new_text": "",
                "rationale": "Remove overbroad catch-all language per guidelines",
                "section": "1. Definition of Confidential Information",
                "priority": "medium",
                "status": "pending",
                "applied": 0,
            },
            {
                "change_type": "replacement",
                "original_text": "State of Delaware",
                "new_text": "State of California",
                "rationale": "Prefer California law for receiving party protection",
                "section": "3. Governing Law",
                "priority": "medium",
                "status": "pending",
                "applied": 0,
            },
        ]
        save_changes(doc_id, mock_changes)
        update_document_status(doc_id, "analyzed")
        print(f"Saved {len(mock_changes)} mock changes")

        # 6. Review changes
        r = client.get(f"/api/documents/{doc_id}")
        doc_data = r.get_json()
        changes = doc_data["changes"]
        assert len(changes) == 3
        print(f"Retrieved {len(changes)} changes for review")

        # 7. Approve all changes
        r = client.post(f"/api/documents/{doc_id}/changes/approve-all")
        assert r.status_code == 200
        print("Approved all changes")

        # 8. Generate markup document
        r = client.post(f"/api/documents/{doc_id}/generate")
        assert r.status_code == 200
        result = r.get_json()
        print(f"Generate result: status={result['status']}")
        print(f"  Stats: {result['stats']}")
        assert result["stats"]["applied"] == 3
        assert result["stats"]["failed"] == 0

        # 9. Download the document
        r = client.get(f"/api/documents/{doc_id}/download")
        assert r.status_code == 200
        assert r.content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        print(f"Download: {len(r.data)} bytes")

        # 10. Verify the downloaded document
        doc_buf = io.BytesIO(r.data)
        result_doc = Document(doc_buf)

        # Check that formatting was applied
        found_strikethrough = False
        found_underline = False
        for para in result_doc.paragraphs:
            for run in para.runs:
                if run.font.strike:
                    found_strikethrough = True
                if run.font.underline:
                    found_underline = True

        assert found_strikethrough, "No strikethrough formatting found in output"
        assert found_underline, "No underline formatting found in output"
        print("Verified: document contains strikethrough and underline formatting")

        print("\n=== Integration test PASSED ===")


if __name__ == "__main__":
    test_full_pipeline()
